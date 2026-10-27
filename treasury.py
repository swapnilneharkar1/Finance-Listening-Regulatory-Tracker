"""
Treasury Peer Benchmarking: reads the NBFC benchmarking Excel and turns it into
clean, chart-ready data for the portal's "Treasury Peer Benchmarking" dashboards.

Workbook conventions it relies on (matching the current NBFCs_Benchmarking.xlsx):
- One visible sheet per company with a borrowing table whose header row has an
  "Instruments" column followed by FY columns (FY20, FY21, ...). The first set of FY
  columns holds amounts (Rs crore); a later repeated set (percentages) is ignored and
  recomputed so every company is measured the same way.
- Below it: "Finance Cost", "Average Borrowing", "Average Borrowing Cost (%)" rows.
- An investment table whose header has "Instrument / bucket" + FY columns, ending at
  "Total Investments" / "Total Gross Investments".
- An "Area | Key readout" table with commentary.
- A summary sheet with "Metric (FYxx) | BFL FYxx | Peer Avg FYxx | Gap vs Peer Avg | Remarks".
- Optional: a "Topic | Explanation | Note" table (e.g. the STRIPS study).
New peer sheets (e.g. Tata Capital, HDB) that follow the same layout are picked up
automatically.
"""
from __future__ import annotations

import re
from pathlib import Path

import openpyxl

# Known company sheets: display name, short label. Order = colour order on charts.
KNOWN_COMPANIES = [
    ("bajaj", "Bajaj Finance", "BFL"),
    ("mahindra", "Mahindra Finance", "Mahindra"),
    ("shriram", "Shriram Finance", "Shriram"),
    ("chola", "Cholamandalam Finance", "Chola"),
    ("tata", "Tata Capital", "Tata Capital"),
    ("hdb", "HDB Financial Services", "HDB"),
]

BORROWING_CATEGORIES = [
    "Secured NCDs", "Unsecured NCDs", "Commercial paper", "Bank term loans",
    "FI & NHB loans", "Cash credit & WCDL", "ECB loans", "ECB bonds",
    "Repo, TREPS & CBLO", "Securitisation", "Deposits", "Subordinated debt", "Other borrowings",
]
INVESTMENT_CATEGORIES = [
    "Government securities", "Treasury bills", "STRIPS", "Mutual funds",
    "Certificates of deposit", "Commercial paper", "PTCs", "Bonds & NCDs",
    "TREPS & repo", "Fixed deposits", "Subsidiaries, associates & JVs", "Equity",
    "Other investments",
]

FY_RE = re.compile(r"^\s*FY\s*'?(\d{2,4})\s*$", re.I)


def _fy(v) -> str | None:
    if not isinstance(v, str):
        return None
    m = FY_RE.match(v)
    if not m:
        return None
    y = m.group(1)
    return f"FY{y[-2:]}"


def _num(v):
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str):
        s = v.strip().replace(",", "").replace("₹", "")
        try:
            return float(s)
        except ValueError:
            return None
    return None


def _text(v) -> str:
    return re.sub(r"\s+", " ", str(v)).strip() if v is not None else ""


def borrowing_category(label: str, section: str) -> str:
    s = label.lower()
    sec = section.lower()
    if "subordinat" in s or "subordinat" in sec:
        return "Subordinated debt"
    if re.search(r"external commercial bond|senior secured note|u\.?s\.? dollar|offshore bond", s):
        return "ECB bonds"
    if "external commercial borrowing" in s or re.search(r"\becb\b", s):
        return "ECB loans"
    if "commercial paper" in s:
        return "Commercial paper"
    if "unsecured" in s and ("ncd" in s or "debenture" in s):
        return "Unsecured NCDs"
    if "secured" in s and ("ncd" in s or "debenture" in s):
        return "Secured NCDs"
    if re.search(r"securiti[sz]ation", s):
        return "Securitisation"
    if re.search(r"term loans? from banks?", s):
        return "Bank term loans"
    if re.search(r"financial institution|national housing bank|\bnhb\b|refinance", s):
        return "FI & NHB loans"
    if re.search(r"cash credit|working capital|wcdl", s):
        return "Cash credit & WCDL"
    if re.search(r"\brepo\b|repurchase|treps|cblo|csgl|tri-?party", s):
        return "Repo, TREPS & CBLO"
    if "deposit" in s or "deposit" in sec:
        return "Deposits"
    return "Other borrowings"


def investment_category(label: str) -> str:
    s = label.lower()
    if "interest accrued" in s or "electoral" in s:
        return "Other investments"
    if "treasury bill" in s or re.search(r"\bt-?bills?\b", s):
        return "Treasury bills"
    if "strips" in s:
        return "STRIPS"
    if "government sec" in s or re.search(r"\bg-?secs?\b", s):
        return "Government securities"
    if "mutual fund" in s:
        return "Mutual funds"
    if re.search(r"certificates? of deposit", s):
        return "Certificates of deposit"
    if "commercial paper" in s:
        return "Commercial paper"
    if "pass through" in s or re.search(r"\bptcs?\b", s):
        return "PTCs"
    if re.search(r"subsidiar|associate|joint venture", s):
        return "Subsidiaries, associates & JVs"
    if re.search(r"treps|tri-?party|\brepo\b", s):
        return "TREPS & repo"
    if "fixed deposit" in s:
        return "Fixed deposits"
    if re.search(r"equity|ccps|shares", s):
        return "Equity"
    if re.search(r"bond|ncd|debenture|debt securit", s):
        return "Bonds & NCDs"
    return "Other investments"


class Sheet:
    def __init__(self, ws):
        self.title = ws.title
        self.rows = [[c for c in r] for r in ws.iter_rows(values_only=True)]
        self.ncols = max((len(r) for r in self.rows), default=0)

    def cell(self, r, c):
        if r < 0 or r >= len(self.rows):
            return None
        row = self.rows[r]
        return row[c] if c < len(row) else None

    def row_texts(self, r):
        return [(c, _text(v)) for c, v in enumerate(self.rows[r]) if _text(v)]

    def find_header(self, pattern: str, start: int = 0):
        """Row index + column of a header cell matching pattern that also has FY columns."""
        rx = re.compile(pattern, re.I)
        for r in range(start, len(self.rows)):
            cells = self.rows[r]
            label_col = next((c for c, v in enumerate(cells) if isinstance(v, str) and rx.search(v)), None)
            if label_col is None:
                continue
            fys = {}
            for c, v in enumerate(cells):
                f = _fy(v)
                if f and c > label_col and f not in fys:  # first occurrence = amounts
                    fys[f] = c
            if len(fys) >= 2:
                return r, label_col, fys
        return None


def _company_meta(sheet_title: str, banner: str):
    key = sheet_title.lower()
    for k, name, short in KNOWN_COMPANIES:
        if k in key or k in banner.lower():
            return name, short
    return sheet_title.strip(), sheet_title.strip()


def parse_company(sh: Sheet) -> dict | None:
    hdr = sh.find_header(r"^\s*instruments?\s*$")
    if not hdr:
        return None
    hr, lc, fys = hdr
    years = sorted(fys, key=lambda f: int(f[2:]))
    banner = " ".join(t for _, t in sh.row_texts(0)) if sh.rows else ""
    name, short = _company_meta(sh.title, banner)
    section_col = lc - 1 if lc > 0 and _text(sh.cell(hr, lc - 1)).lower().startswith("section") else None

    # --- Borrowing table ---
    leaves, totals, grand_total = [], [], None
    extra = {}  # finance cost, average borrowing, cost
    r = hr + 1
    blanks = 0
    while r < len(sh.rows):
        label = _text(sh.cell(r, lc))
        if not label:
            # rows with only numbers (helper rows) or blank
            if not any(_text(v) for v in sh.rows[r]):
                blanks += 1
                if blanks >= 3:
                    break
            r += 1
            continue
        blanks = 0
        if re.match(r"^\s*instrument\s*/\s*bucket", label, re.I):
            break
        vals = {f: _num(sh.cell(r, c)) for f, c in fys.items()}
        low = label.lower()
        section = _text(sh.cell(r, section_col)) if section_col is not None else ""
        if low.startswith("finance cost"):
            extra["finance_cost"] = vals
        elif low.startswith("average borrowing cost") and "ecb" not in low:
            extra["avg_cost"] = vals
        elif low.startswith("average borrowing") and "ecb" in low:
            extra["avg_cost_ecb"] = vals
        elif low.startswith("average borrowing"):
            extra["avg_borrowing"] = vals
        elif low.rstrip() == "total":
            grand_total = vals
        elif low.startswith("total"):
            totals.append({"label": label, "values": vals})
        elif not extra:  # still inside the instrument table
            leaves.append({"label": label, "section": section, "values": vals})
        r += 1

    for leaf in leaves:
        leaf["category"] = borrowing_category(leaf["label"], leaf["section"])
    _infer_sections_from_totals(sh, hr, lc, leaves)

    if grand_total is None:
        grand_total = {f: sum((l["values"][f] or 0) for l in leaves) for f in years}

    borrow_mix = _mix(leaves, years, grand_total, BORROWING_CATEGORIES)

    # Average cost: use the sheet's own figure; fall back to finance cost / avg borrowing.
    avg_cost = extra.get("avg_cost") or {}
    for f in years:
        if avg_cost.get(f) is None and extra.get("finance_cost") and extra.get("avg_borrowing"):
            fc, ab = extra["finance_cost"].get(f), extra["avg_borrowing"].get(f)
            if fc and ab:
                avg_cost[f] = fc / ab
        if avg_cost.get(f) is not None and avg_cost[f] > 1:  # entered as 7.5 instead of 7.5%
            avg_cost[f] = avg_cost[f] / 100

    # --- Investment table ---
    inv = None
    ih = sh.find_header(r"^\s*instrument\s*/\s*bucket", hr + 1)
    if ih:
        ir, ilc, ifys = ih
        iyears = sorted(ifys, key=lambda f: int(f[2:]))
        ileaves, itotal = [], None
        for r in range(ir + 1, len(sh.rows)):
            label = _text(sh.cell(r, ilc))
            if not label:
                if not any(_text(v) for v in sh.rows[r]):
                    break
                continue
            low = label.lower()
            vals = {f: _num(sh.cell(r, c)) for f, c in ifys.items()}
            if low.startswith("less") or low.startswith("total net"):
                continue
            if low.startswith("total"):
                itotal = vals
                if "gross" in low or low.strip() == "total investments":
                    break
                continue
            ileaves.append({"label": label, "values": vals, "category": investment_category(label)})
        if itotal is None:
            itotal = {f: sum((l["values"][f] or 0) for l in ileaves) for f in iyears}
        inv = {
            "years": iyears,
            "rows": [{"label": l["label"], "category": l["category"],
                      "values": l["values"]} for l in ileaves],
            "total": itotal,
            "mix": _mix(ileaves, iyears, itotal, INVESTMENT_CATEGORIES),
        }

    return {
        "sheet": sh.title,
        "name": name,
        "short": short,
        "borrowing": {
            "years": years,
            "rows": [{"label": l["label"], "category": l["category"], "values": l["values"]} for l in leaves],
            "subtotals": totals,
            "total": grand_total,
            "mix": borrow_mix,
            "finance_cost": extra.get("finance_cost", {}),
            "avg_borrowing": extra.get("avg_borrowing", {}),
            "avg_cost": avg_cost,
        },
        "investments": inv,
        "readouts": parse_readouts(sh),
    }


def _infer_sections_from_totals(sh: Sheet, hr: int, lc: int, leaves: list):
    """For sheets without a Section column: a leaf belongs to the next 'Total X' row below it."""
    if any(l["section"] for l in leaves):
        return
    order = []
    for r in range(hr + 1, len(sh.rows)):
        label = _text(sh.cell(r, lc))
        if label:
            order.append(label)
    pending = []
    for label in order:
        low = label.lower()
        if low.startswith("total"):
            for p in pending:
                if "deposit" in low and p["category"] == "Other borrowings":
                    p["category"] = "Deposits"
                if "subordinat" in low and p["category"] == "Other borrowings":
                    p["category"] = "Subordinated debt"
            pending = []
            if low.strip() == "total":
                break
            continue
        leaf = next((l for l in leaves if l["label"] == label and l not in pending), None)
        if leaf:
            pending.append(leaf)


def _mix(leaves, years, total, categories):
    """Category -> {FY: share of total}, plus amounts."""
    amounts = {c: {f: 0.0 for f in years} for c in categories}
    for l in leaves:
        for f in years:
            v = l["values"].get(f)
            if v:
                amounts[l["category"]][f] += v
    shares = {}
    for c in categories:
        shares[c] = {}
        for f in years:
            t = total.get(f) if total else None
            shares[c][f] = (amounts[c][f] / t) if t else None
    used = [c for c in categories if any(amounts[c][f] for f in years)]
    return {"categories": used, "amounts": {c: amounts[c] for c in used},
            "shares": {c: shares[c] for c in used}}


def parse_readouts(sh: Sheet) -> list:
    out = []
    for r in range(len(sh.rows)):
        cells = sh.row_texts(r)
        if len(cells) >= 2 and cells[0][1].lower() == "area" and "key readout" in cells[1][1].lower():
            c0, c1 = cells[0][0], cells[1][0]
            for rr in range(r + 1, len(sh.rows)):
                a, b = _text(sh.cell(rr, c0)), _text(sh.cell(rr, c1))
                if not a and not b:
                    break
                out.append({"area": a, "text": b})
            break
    return out


def parse_summary(sh: Sheet) -> list:
    """Tables shaped 'Metric (FYxx) | BFL | Peer Avg | Gap | Remarks' with a title row above."""
    tables = []
    for r in range(len(sh.rows)):
        cells = sh.row_texts(r)
        if not cells or not cells[0][1].lower().startswith("metric"):
            continue
        heads = [t for _, t in cells]
        if not any("peer" in h.lower() for h in heads):
            continue
        cols = [c for c, _ in cells]
        title = " ".join(t for _, t in sh.row_texts(r - 1)[:1]) if r > 0 else ""
        fy = _fy((re.search(r"FY\s*\d{2,4}", heads[0]) or [""])[0] or "") or ""
        rows = []
        for rr in range(r + 1, len(sh.rows)):
            metric = _text(sh.cell(rr, cols[0]))
            if not metric:
                break
            vals = [_num(sh.cell(rr, c)) for c in cols[1:4]]
            remark = _text(sh.cell(rr, cols[4])) if len(cols) > 4 else ""
            rows.append({"metric": metric, "bfl": vals[0], "peer": vals[1],
                         "gap": vals[2] if vals[2] is not None else (
                             (vals[0] - vals[1]) if None not in vals[:2] else None),
                         "remark": remark})
        kind = "investment" if "invest" in title.lower() else "borrowing"
        tables.append({"title": title.strip() or heads[0], "kind": kind, "fy": fy,
                       "columns": heads[:5], "rows": rows})
    return tables


def parse_topic_tables(sh: Sheet) -> list:
    """'Topic | Explanation | Note' study tables (e.g. STRIPS), with the question list above."""
    out = []
    for r in range(len(sh.rows)):
        cells = sh.row_texts(r)
        heads = [t.lower() for _, t in cells]
        if "topic" in heads and "explanation" in heads:
            cols = [c for c, _ in cells]
            rows = []
            for rr in range(r + 1, len(sh.rows)):
                vals = [_text(sh.cell(rr, c)) for c in cols]
                if not any(vals):
                    break
                rows.append(vals)
            # title + questions: look upwards for a "Study of ..." label
            title, questions = "", []
            for up in range(r - 1, max(-1, r - 12), -1):
                for c, t in sh.row_texts(up):
                    if t.lower().startswith("study of"):
                        title = t.strip()
                        qs = []
                        for q in range(up, r):
                            for cc, tt in sh.row_texts(q):
                                if cc > c and tt:
                                    qs.append(tt)
                        questions = qs
                        break
                if title:
                    break
            out.append({"title": title or "Study", "questions": questions,
                        "columns": [t for _, t in cells], "rows": rows})
    return out


def load_workbook_data(path: str | Path, include_hidden: bool = False) -> dict:
    wb = openpyxl.load_workbook(path, data_only=True, read_only=False)
    companies, summary, studies, notes = [], [], [], []
    for ws in wb.worksheets:
        if ws.sheet_state != "visible" and not include_hidden:
            continue
        sh = Sheet(ws)
        comp = parse_company(sh)
        if comp:
            companies.append(comp)
            continue
        s = parse_summary(sh)
        if s:
            summary.extend(s)
        studies.extend(parse_topic_tables(sh))
    # colour order = KNOWN_COMPANIES order, then any new sheets
    order = {name: i for i, (_, name, _) in enumerate(KNOWN_COMPANIES)}
    companies.sort(key=lambda c: order.get(c["name"], 99))
    all_years = sorted({f for c in companies for f in c["borrowing"]["years"]}, key=lambda f: int(f[2:]))
    common = [f for f in all_years if all(f in c["borrowing"]["years"] for c in companies)]
    return {
        "source": Path(path).name,
        "companies": companies,
        "summary": summary,
        "studies": studies,
        "years": all_years,
        "default_year": common[-1] if common else (all_years[-1] if all_years else ""),
        "borrowing_categories": BORROWING_CATEGORIES,
        "investment_categories": INVESTMENT_CATEGORIES,
    }
